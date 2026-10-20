"""Configured MBOM input template: create it, read it, and check it.

The workbook has two working tabs:
  Model      model version, the item the model is attached to, the configured
             item, the variant and its values (one row per X column)
  Structure  one row per BOM line, top-down, with X marks per variant value

Everything the builder needs is recomputed here in Python, so the Excel
formulas (Effectivity / Check columns) are a convenience for the person filling
the sheet, never something the tool relies on.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ITEM_TYPES = ["CreateAssembly", "Provide", "CreateMaterial", "CreateKit"]
VALUE_COLUMNS = "LMNOPQ"          # up to 6 variant values
FIRST_ROW = 5                     # first data row on the Structure tab
HEADER_ROW = 4
FORMULA_ROWS = 300                # rows pre-filled with formulas / dropdowns


# --------------------------------------------------------------------------
# Reading and checking
# --------------------------------------------------------------------------
@dataclass
class Line:
    row: int
    level: int
    part_number: str
    title: str
    description: str
    type: str
    uom: str
    qty: int
    qty_note: str
    source: str            # New | Existing
    existing_id: str
    values: list[str]      # variant value names this line belongs to (from X marks)
    parent: int | None = None   # index into lines
    effectivity: str = ""       # "", "ALL" or "1143, 1144"
    problems: list[str] = field(default_factory=list)


@dataclass
class Plan:
    file: str
    model_title: str
    model_revision: str
    attach_to: str
    configured_item: str
    variant: str
    values: list[dict]            # [{name, part_number, description}]
    lines: list[Line]
    problems: list[str]

    @property
    def ok(self) -> bool:
        return not self.problems and all(not ln.problems for ln in self.lines)

    def index_of(self, title: str) -> int | None:
        for i, ln in enumerate(self.lines):
            if ln.title == title:
                return i
        return None

    def children(self, idx: int) -> list[int]:
        return [i for i, ln in enumerate(self.lines) if ln.parent == idx]

    def path_down(self) -> list[int]:
        """Indexes from the attach-to item down to the configured item."""
        a, c = self.index_of(self.attach_to), self.index_of(self.configured_item)
        if a is None or c is None:
            return []
        chain, i = [], c
        while i is not None:
            chain.append(i)
            if i == a:
                return list(reversed(chain))
            i = self.lines[i].parent
        return []

    def summary(self) -> dict:
        c = self.index_of(self.configured_item)
        kids = self.children(c) if c is not None else []
        out = {}
        for v in self.values:
            n = [i for i in kids if self.lines[i].effectivity == "ALL" or v["name"] in self.lines[i].values]
            out[v["name"]] = {"lines": len(n), "instances": sum(self.lines[i].qty for i in n)}
        return {
            "lines": len(self.lines),
            "instances": sum(ln.qty for ln in self.lines if ln.parent is not None),
            "new_items": sum(1 for ln in self.lines if ln.source == "New"),
            "existing_items": sum(1 for ln in self.lines if ln.source == "Existing"),
            "common_lines": sum(1 for i in kids if self.lines[i].effectivity == "ALL"),
            "variant_specific_lines": sum(1 for i in kids if self.lines[i].effectivity not in ("", "ALL")),
            "per_value": out,
        }


def _s(v) -> str:
    return "" if v is None else " ".join(str(v).split())


def read_template(path: str | Path) -> Plan:
    wb = load_workbook(path, data_only=False)
    problems: list[str] = []
    for tab in ("Model", "Structure"):
        if tab not in wb.sheetnames:
            raise ValueError(f"'{Path(path).name}' has no '{tab}' tab. Start from the Configured MBOM template.")
    m, s = wb["Model"], wb["Structure"]

    model_title, model_rev = _s(m["B3"].value), _s(m["B4"].value) or "A"
    attach_to, configured, variant = _s(m["B5"].value), _s(m["B6"].value), _s(m["B7"].value)
    for label, val in (("Model version title (Model!B3)", model_title), ("Attach model to item (Model!B5)", attach_to),
                       ("Configured item (Model!B6)", configured), ("Variant name (Model!B7)", variant)):
        if not val:
            problems.append(f"{label} is empty")

    values = []
    for k in range(len(VALUE_COLUMNS)):
        r = 12 + k
        name = _s(m.cell(r, 2).value)
        if name:
            values.append({"name": name, "part_number": _s(m.cell(r, 3).value),
                           "description": _s(m.cell(r, 4).value), "column": VALUE_COLUMNS[k]})
    if not values:
        problems.append("No variant values on the Model tab (rows 12-17, column B)")
    if len({v["name"] for v in values}) != len(values):
        problems.append("Variant value names on the Model tab must be unique")

    lines: list[Line] = []
    r = FIRST_ROW
    while r <= s.max_row:
        lvl = s.cell(r, 2).value
        title = _s(s.cell(r, 4).value) or _s(s.cell(r, 3).value)
        if lvl in (None, "") and not title:
            r += 1
            continue
        ln = Line(row=r, level=-1, part_number=_s(s.cell(r, 3).value), title=title,
                  description=_s(s.cell(r, 5).value), type=_s(s.cell(r, 6).value) or "Provide",
                  uom=_s(s.cell(r, 7).value) or "Each", qty=0, qty_note=_s(s.cell(r, 9).value),
                  source=_s(s.cell(r, 10).value) or "New", existing_id=_s(s.cell(r, 11).value), values=[])
        try:
            ln.level = int(lvl)
        except (TypeError, ValueError):
            ln.problems.append("Level is not a whole number")
        q = s.cell(r, 8).value
        try:
            ln.qty = int(float(q)) if q not in (None, "") else 0
            if q not in (None, "") and float(q) != int(float(q)):
                ln.problems.append("Qty must be a whole number")
        except (TypeError, ValueError):
            ln.problems.append("Qty is not a number")
        if not title:
            ln.problems.append("Title is empty")
        if ln.type not in ITEM_TYPES:
            ln.problems.append(f"Type '{ln.type}' is not one of {', '.join(ITEM_TYPES)}")
        if ln.source not in ("New", "Existing"):
            ln.problems.append("Item source must be New or Existing")
        if ln.source == "Existing" and not ln.existing_id:
            ln.problems.append("Existing ID missing")
        for v in values:
            mark = _s(s[f"{v['column']}{r}"].value).upper()
            if mark == "X":
                ln.values.append(v["name"])
            elif mark:
                ln.problems.append(f"Unexpected mark '{mark}' in column {v['column']} (use X)")
        lines.append(ln)
        r += 1

    if not lines:
        problems.append("The Structure tab has no lines")

    # parents and level checks
    stack: list[int] = []
    for i, ln in enumerate(lines):
        if ln.level < 0:
            continue
        if i == 0 and ln.level != 0:
            ln.problems.append("The first line must be level 0 (the top item)")
        if i > 0 and ln.level == 0:
            ln.problems.append("Only one level-0 line is allowed")
        prev_level = lines[i - 1].level if i > 0 else -1
        if i > 0 and ln.level > prev_level + 1:
            ln.problems.append("Level skips a step")
        while stack and lines[stack[-1]].level >= ln.level:
            stack.pop()
        ln.parent = stack[-1] if stack else None
        stack.append(i)
        if ln.parent is not None and ln.qty <= 0:
            ln.problems.append("Qty missing or 0 (use 1 with a Qty note for consumables)")

    # children must sit under assemblies
    for i, ln in enumerate(lines):
        if ln.parent is not None and lines[ln.parent].type == "Provide":
            lines[ln.parent].problems.append("Has child lines but Type is Provide (use CreateAssembly)")

    titles = [ln.title for ln in lines if ln.source == "New"]
    dups = {t for t in titles if titles.count(t) > 1}
    for ln in lines:
        if ln.title in dups and ln.source == "New":
            ln.problems.append("Same title appears twice as New; list it once and reuse it, or mark the repeat Existing")

    plan = Plan(str(path), model_title, model_rev, attach_to, configured, variant, values, lines, problems)

    ci = plan.index_of(configured)
    if configured and ci is None:
        problems.append(f"Configured item '{configured}' is not a Title on the Structure tab")
    if attach_to and plan.index_of(attach_to) is None:
        problems.append(f"Attach-to item '{attach_to}' is not a Title on the Structure tab")
    if ci is not None and attach_to and plan.index_of(attach_to) is not None and not plan.path_down():
        problems.append("The configured item is not below the attach-to item")
    for i in plan.path_down():
        if lines[i].type == "Provide":
            lines[i].problems.append("Lies on the configured path but is a Provide")

    kids = set(plan.children(ci)) if ci is not None else set()
    for i, ln in enumerate(lines):
        if i in kids:
            if not ln.values:
                ln.problems.append("Not in any variant (put X under at least one value)")
            ln.effectivity = "ALL" if len(ln.values) == len(values) else ", ".join(ln.values)
        elif ln.values:
            ln.problems.append("X marks outside the configured item's direct children")
    return plan


def plan_report(plan: Plan) -> dict:
    """JSON-friendly summary for the preview tool."""
    rows = [{
        "row": ln.row, "level": ln.level, "title": ln.title, "type": ln.type, "qty": ln.qty,
        "source": ln.source, "parent": plan.lines[ln.parent].title if ln.parent is not None else "",
        "effectivity": ln.effectivity, "problems": ln.problems,
    } for ln in plan.lines]
    return {
        "file": plan.file, "ok": plan.ok,
        "model": {"title": plan.model_title, "revision": plan.model_revision, "variant": plan.variant,
                  "values": [v["name"] for v in plan.values],
                  "attach_to": plan.attach_to, "configured_item": plan.configured_item,
                  "attached_items": [plan.lines[i].title for i in plan.path_down()]},
        "summary": plan.summary(),
        "problems": plan.problems + [f"row {ln.row} ({ln.title}): {p}" for ln in plan.lines for p in ln.problems],
        "lines": rows,
    }


# --------------------------------------------------------------------------
# Creating a template (blank or pre-filled)
# --------------------------------------------------------------------------
def make_template(path: str | Path, model: dict | None = None, values: list[dict] | None = None,
                  rows: list[dict] | None = None) -> Path:
    """Write a template. `rows` items: level, part_number, title, description,
    type, uom, qty, qty_note, source, existing_id, marks (list of value names)."""
    model = model or {}
    values = values or []
    rows = rows or []
    F, NAVY = "Arial", "0B3D6B"
    INPUT, CALC = PatternFill("solid", fgColor="FFF7D6"), PatternFill("solid", fgColor="EEF1F5")
    HDR, VHDR = PatternFill("solid", fgColor=NAVY), PatternFill("solid", fgColor="2E6DA4")
    hfont, body, bold = Font(name=F, bold=True, color="FFFFFF", size=10), Font(name=F, size=10), Font(name=F, size=10, bold=True)
    small = Font(name=F, size=9, italic=True, color="5A6878")
    thin = Side(style="thin", color="C9D3DE")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap, top, center = Alignment(wrap_text=True, vertical="top"), Alignment(vertical="top"), Alignment(horizontal="center", vertical="top")

    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width, ins.column_dimensions["B"].width, ins.column_dimensions["C"].width = 3, 24, 95
    ins["B2"] = "Configured MBOM input template"
    ins["B2"].font = Font(name=F, size=16, bold=True, color=NAVY)
    ins["B3"] = "One workbook per configured product. Built by the Configured MBOM Builder (preview first, then build)."
    ins["B3"].font = small
    guide = [
        ("How it is used", "Ask Claude to preview the workbook. The preview lists exactly what will be created and any problems. Build only when the preview is clean."),
        ("Yellow cells", "Input. Grey cells are calculated; do not type in them."),
        ("Model tab", "Model version title and revision, the item the model is attached to (usually the top item), the configured item (whose direct children vary), the variant name, and one row per variant value (Description is optional: use only an official description, never a list of BOM differences, which goes stale). Row 12 is the first X column (L) on the Structure tab, row 13 the second (M), and so on."),
        ("Structure tab", "One row per BOM line, top-down. Level 0 is the top item. A line's parent is the closest row above it with a lower level."),
        ("Title / Part number", "Title is what the Manufacturing Item is called in 3DEXPERIENCE. Part number is for reference (often the same)."),
        ("Type", "CreateAssembly for anything with children, Provide for purchased or leaf parts."),
        ("Qty", "Whole units. Quantities above 1 become that many instances. For consumables use 1 and write 'As required' in Qty note."),
        ("Item source", "New creates the item. Existing reuses an item already in 3DEXPERIENCE: put its physical ID in Existing ID."),
        ("Variant X columns", "Only on direct children of the configured item. X under every value the line belongs to. X in every column means a common part (no effectivity)."),
        ("Effectivity / Check", "Effectivity shows what will be applied (ALL, or the values). Check must say OK on every row."),
        ("Rerunning", "A model version title that already exists stops the build, so a workbook cannot be built twice by accident. Change the title to build again."),
    ]
    for i, (a, b) in enumerate(guide):
        ins.cell(5 + i, 2, a).font = bold
        c = ins.cell(5 + i, 3, b); c.font, c.alignment = body, wrap
    r0 = 6 + len(guide)
    ins.cell(r0, 2, "Legend").font = bold
    for j, (t, f) in enumerate((("Input", INPUT), ("Calculated", CALC))):
        c = ins.cell(r0 + 1 + j, 2, t); c.fill, c.font, c.border = f, body, box

    m = wb.create_sheet("Model")
    for col, w in zip("ABCD", (30, 22, 24, 70)):
        m.column_dimensions[col].width = w
    m["A1"] = "Model and variants"; m["A1"].font = Font(name=F, size=14, bold=True, color=NAVY)
    fields = [("Model version title", model.get("title"), "Portfolio model version to create. Must not exist yet."),
              ("Model version revision", model.get("revision", "A"), ""),
              ("Attach model to item", model.get("attach_to"), "Title of the item the model is attached to (usually the level-0 item). Every assembly down to the configured item gets it too."),
              ("Configured item", model.get("configured_item"), "Title of the item whose direct children carry variant effectivity."),
              ("Variant name", model.get("variant"), "Portfolio variant to create.")]
    for i, (a, b, c) in enumerate(fields):
        rr = 3 + i
        m.cell(rr, 1, a).font = bold
        cb = m.cell(rr, 2, b); cb.fill, cb.font, cb.border = INPUT, body, box
        m.cell(rr, 3, c).font = small
        m.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=4)
    m["A8"] = "Configured item level"; m["A8"].font = bold
    m["B8"] = '=IFERROR(INDEX(Structure!$B:$B,MATCH($B$6,Structure!$D:$D,0)),"not found")'
    m["B8"].fill, m["B8"].font, m["B8"].border = CALC, body, box
    m["A10"] = "Variant values"; m["A10"].font = Font(name=F, size=11, bold=True, color=NAVY)
    for j, h in enumerate(["X column", "Value (short name)", "Full part number", "Description"], 1):
        c = m.cell(11, j, h); c.font, c.fill, c.border = hfont, HDR, box
    for k in range(len(VALUE_COLUMNS)):
        rr = 12 + k
        a = m.cell(rr, 1, f"Structure column {VALUE_COLUMNS[k]}"); a.font, a.border = body, box
        v = values[k] if k < len(values) else {}
        for j, key in ((2, "name"), (3, "part_number"), (4, "description")):
            c = m.cell(rr, j, v.get(key)); c.fill, c.font, c.border, c.alignment = INPUT, body, box, wrap
    m["A19"] = "Summary"; m["A19"].font = Font(name=F, size=11, bold=True, color=NAVY)
    for j, h in enumerate(["Value", "Lines in this variant", "Of which variant-specific"], 1):
        c = m.cell(20, j, h); c.font, c.fill, c.border = hfont, HDR, box
    for k, L in enumerate(VALUE_COLUMNS):
        rr = 21 + k
        m.cell(rr, 1, f'=IF($B${12+k}="","",$B${12+k})')
        m.cell(rr, 2, f'=IF($B${12+k}="","",COUNTIF(Structure!$R:$R,"ALL")+COUNTIFS(Structure!${L}:${L},"X",Structure!$R:$R,"<>ALL"))')
        m.cell(rr, 3, f'=IF($B${12+k}="","",COUNTIFS(Structure!${L}:${L},"X",Structure!$R:$R,"<>ALL"))')
        for j in range(1, 4):
            c = m.cell(rr, j); c.fill, c.font, c.border, c.alignment = CALC, body, box, Alignment(horizontal="left")
    m["A28"] = "Common lines (in every variant)"; m["A28"].font = bold
    m["B28"] = '=COUNTIF(Structure!$R:$R,"ALL")'
    m["A29"] = "Rows with a problem"; m["A29"].font = bold
    m["B29"] = f'=COUNTIFS(Structure!$S${FIRST_ROW}:$S${FORMULA_ROWS},"?*",Structure!$S${FIRST_ROW}:$S${FORMULA_ROWS},"<>OK")'
    m["C29"] = "Must be 0 before building."; m["C29"].font = small
    for a in ("B28", "B29"):
        m[a].fill, m[a].font, m[a].border = CALC, body, box

    s = wb.create_sheet("Structure")
    s["A1"] = "MBOM structure"; s["A1"].font = Font(name=F, size=14, bold=True, color=NAVY)
    s["A2"] = "One row per line, top-down. X marks only on direct children of the configured item (see Model tab)."; s["A2"].font = small
    cols = [("A", "Line", 6), ("B", "Level", 7), ("C", "Part number", 18), ("D", "Title", 18), ("E", "Description", 52),
            ("F", "Type", 16), ("G", "UOM", 8), ("H", "Qty", 6), ("I", "Qty note", 14), ("J", "Item source", 11), ("K", "Existing ID", 16)]
    for col, h, w in cols:
        c = s[f"{col}{HEADER_ROW}"]; c.value, c.font, c.fill, c.border = h, hfont, HDR, box
        c.alignment = Alignment(wrap_text=True, vertical="center")
        s.column_dimensions[col].width = w
    for k, L in enumerate(VALUE_COLUMNS):
        c = s[f"{L}{HEADER_ROW}"]
        c.value = f'=IF(Model!$B${12+k}="","",Model!$B${12+k})'
        c.font, c.fill, c.border, c.alignment = hfont, VHDR, box, Alignment(horizontal="center")
        s.column_dimensions[L].width = 7
    for L, h, w in (("R", "Effectivity", 18), ("S", "Check", 30)):
        c = s[f"{L}{HEADER_ROW}"]; c.value, c.font, c.fill, c.border = h, hfont, HDR, box
        s.column_dimensions[L].width = w
    s["L3"] = "Variant values (X)"; s["L3"].font = Font(name=F, size=9, bold=True, color="2E6DA4")
    s.merge_cells("L3:Q3")

    nv = "COUNTA(Model!$B$12:$B$17)"
    for i in range(FORMULA_ROWS - FIRST_ROW + 1):
        rr = FIRST_ROW + i
        d = rows[i] if i < len(rows) else None
        vals = [None] * 11
        if d:
            vals = [i + 1, d.get("level"), d.get("part_number"), d.get("title"), d.get("description"), d.get("type"),
                    d.get("uom", "Each"), d.get("qty"), d.get("qty_note"), d.get("source", "New"), d.get("existing_id")]
        else:
            vals[0] = f'=IF(B{rr}="","",ROW()-{FIRST_ROW - 1})'
        for j, v in enumerate(vals, 1):
            c = s.cell(rr, j, v); c.font, c.border = body, box
            c.alignment = wrap if j == 5 else top
            c.fill = CALC if j == 1 else INPUT
        for k, L in enumerate(VALUE_COLUMNS):
            mark = "X" if d and k < len(values) and values[k].get("name") in (d.get("marks") or []) else None
            c = s[f"{L}{rr}"]; c.value, c.fill, c.font, c.border, c.alignment = mark, INPUT, bold, box, center
        xc = f'COUNTIF(L{rr}:Q{rr},"X")'
        under = f"B{rr}=Model!$B$8+1"
        s.cell(rr, 18, f'=IF(B{rr}="","",IF({xc}=0,"",IF({xc}={nv},"ALL",_xlfn.TEXTJOIN(", ",TRUE,'
                       + ",".join(f'IF({L}{rr}="X",{L}${HEADER_ROW},"")' for L in VALUE_COLUMNS) + "))))")
        lvl_check = f"B{rr}>B{rr-1}+1" if rr > FIRST_ROW else "FALSE"
        s.cell(rr, 19, f'=IF(B{rr}="","",IF({lvl_check},"Level skips a step",'
                       f'IF(AND(B{rr}>0,OR(H{rr}="",H{rr}<=0)),"Qty missing or 0",'
                       f'IF(AND(J{rr}="Existing",K{rr}=""),"Existing ID missing",'
                       f'IF(AND({under},{xc}=0),"Not in any variant",'
                       f'IF(AND(NOT({under}),{xc}>0),"X outside configured item","OK"))))))')
        for j in (18, 19):
            c = s.cell(rr, j); c.fill, c.font, c.border = CALC, body, box
    s.freeze_panes = f"F{FIRST_ROW}"

    lists = wb.create_sheet("Lists")
    for i, t in enumerate(ITEM_TYPES, 1):
        lists.cell(i, 1, t)
    for i, t in enumerate(["New", "Existing"], 1):
        lists.cell(i, 2, t)
    uoms = ["Each", "g", "kg", "mm", "m", "ml", "l", "in"]
    for i, t in enumerate(uoms, 1):
        lists.cell(i, 3, t)
    lists.sheet_state = "hidden"
    end = FORMULA_ROWS
    for formula, rng in ((f"=Lists!$A$1:$A${len(ITEM_TYPES)}", f"F{FIRST_ROW}:F{end}"),
                         ("=Lists!$B$1:$B$2", f"J{FIRST_ROW}:J{end}"),
                         (f"=Lists!$C$1:$C${len(uoms)}", f"G{FIRST_ROW}:G{end}"),
                         ('"X"', f"L{FIRST_ROW}:Q{end}")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        s.add_data_validation(dv); dv.add(rng)
    dvl = DataValidation(type="whole", operator="between", formula1="0", formula2="20", allow_blank=True)
    s.add_data_validation(dvl); dvl.add(f"B{FIRST_ROW}:B{end}")
    red = PatternFill("solid", fgColor="F8D7DA")
    s.conditional_formatting.add(f"S{FIRST_ROW}:S{end}", FormulaRule(
        formula=[f'AND(S{FIRST_ROW}<>"",S{FIRST_ROW}<>"OK")'], fill=red, font=Font(name=F, color="9C0006", bold=True)))
    s.cell(HEADER_ROW, 8).comment = Comment("Whole units. Quantities above 1 become that many instances.", "Configured MBOM Builder")
    s.cell(HEADER_ROW, 11).comment = Comment("Physical ID of an existing Manufacturing Item; only when Item source = Existing.", "Configured MBOM Builder")
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
